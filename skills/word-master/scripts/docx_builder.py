# -*- coding: utf-8 -*-
"""word-master builder: content.md → templated DOCX (deterministic injection).

The AI authors content grammar only (references/content-grammar.md); every
format decision here comes from the template contract. Never hand-edit the
built docx — fix content.md or the contract and rebuild.

Usage: docx_builder.py build <project_path>
"""
import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu
from docx.text.paragraph import Paragraph
from wm_common import TEMPLATES, WMError, load_contract, load_project, resolve_project, save_project

FACT_RE = re.compile(r"<!--\s*(F\d+)\s*-->")
REF_RE = re.compile(r"\{\{ref:(fig|tbl):([\w\-]+)\}\}")
MAX_IMG_WIDTH = Cm(14.5)  # A4 minus contract margins, safety-capped


# ---------------------------------------------------------------- XML helpers
def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def para_with_style(doc, style_id, outline_lvl=None):
    """Append a paragraph with explicit pStyle (styleId). Returns Paragraph."""
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    ppr.append(el("w:pStyle", val=style_id))
    if outline_lvl is not None:
        ppr.append(el("w:outlineLvl", val=outline_lvl))
    p.append(ppr)
    doc.element.body.insert(list(doc.element.body).index(doc.element.body.find(qn("w:sectPr"))), p)
    return Paragraph(p, doc._body)


def add_text_run(para, text, bold=False, italic=False):
    r = para.add_run(text)
    r.bold = bold or None
    r.italic = italic or None
    return r


def add_field(para, instr, dirty=True):
    """Insert a field (begin/instr/separate/end) with empty result."""
    r1 = el("w:r")
    fc = el("w:fldChar", fldCharType="begin")
    if dirty:
        fc.set(qn("w:dirty"), "true")
    r1.append(fc)
    r2 = el("w:r")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2.append(it)
    r3 = el("w:r")
    r3.append(el("w:fldChar", fldCharType="separate"))
    r4 = el("w:r")
    r4.append(el("w:fldChar", fldCharType="end"))
    for r in (r1, r2, r3, r4):
        para._p.append(r)


class Bookmarks:
    def __init__(self):
        self._id = 100

    def wrap(self, para, name):
        bid = str(self._id)
        self._id += 1
        start = el("w:bookmarkStart", id=bid, name=name)
        end = el("w:bookmarkEnd", id=bid)
        ppr = para._p.find(qn("w:pPr"))
        idx = list(para._p).index(ppr) + 1 if ppr is not None else 0
        para._p.insert(idx, start)
        para._p.append(end)


# ---------------------------------------------------------------- content parser
def parse_content(md_text):
    """Parse content grammar into an ordered item list.

    Items: ('title', text) ('heading', level, text) ('para', text)
           ('bullet', text) ('figure', slug, img_path)
           ('caption_fig', slug, text) ('caption_tbl', slug, text)
           ('table', rows) — rows: list of list of str, first row = header
    """
    items = []
    lines = md_text.splitlines()
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("@title:"):
            items.append(("title", s[len("@title:"):].strip()))
            i += 1
            continue
        if s.startswith("@template:"):
            i += 1
            continue
        if s.startswith(":::caption"):
            m = re.match(r":::caption\s+(fig|tbl)\s+([\w\-]+)\s+(.*)", s)
            if not m:
                raise WMError(f"caption 语法错误（第{i + 1}行）: {s}")
            kind, slug, text = m.group(1), m.group(2), m.group(3).strip()
            if kind == "tbl":
                # table body follows: blank, pipe table, blank, ":::"
                i += 1
                rows = []
                while i < n:
                    t = lines[i].strip()
                    if t == ":::":
                        break
                    if t.startswith("|"):
                        cells = [c.strip() for c in t.strip("|").split("|")]
                        rows.append(cells)
                    i += 1
                else:
                    raise WMError("表格块未闭合（缺 :::）")
                if rows and set(rows[1]) <= {"---", ":--", "--:", ":-:"}:
                    rows.pop(1)
                items.append(("caption_tbl", slug, text))
                items.append(("table", rows))
                i += 1
                continue
            items.append(("caption_fig", slug, text))
            i += 1
            continue
        if s.startswith(":::figure"):
            slug = s[len(":::figure"):].strip()
            img = None
            i += 1
            while i < n:
                t = lines[i].strip()
                if t == ":::":
                    break
                if t.startswith("<img:"):
                    img = t[len("<img:"):].strip().rstrip(">")
                i += 1
            else:
                raise WMError("figure 块未闭合（缺 :::）")
            if not img:
                raise WMError(f"figure {slug} 缺 <img: 路径>")
            items.append(("figure", slug, img))
            i += 1
            continue
        if re.match(r"^#{8,}\s+", s):
            # content-grammar: ######## and deeper is a syntax error, not body text
            raise WMError(f"标题级别超过 7 级（第{i + 1}行）: {s.split()[0]}")
        m = re.match(r"^(#{1,7})\s+(.*)$", s)
        if m:
            items.append(("heading", len(m.group(1)), m.group(2).strip()))
            i += 1
            continue
        if s.startswith("- "):
            items.append(("bullet", s[2:].strip()))
            i += 1
            continue
        if s.startswith(":::"):
            raise WMError(f"未识别的块语法（第{i + 1}行）: {s}")
        items.append(("para", s))
        i += 1
    return items


# ---------------------------------------------------------------- renderer
class Renderer:
    def __init__(self, doc, contract, bookmarks, project_dir=None):
        self.project_dir = Path(project_dir) if project_dir else Path.cwd()
        self.doc = doc
        self.c = contract
        self.bm = bookmarks
        self.styles = contract["styles"]
        self.facts_cited = []

    def toc_page(self, depth):
        para_with_style(self.doc, self.styles["toc_heading"]).add_run("目  录")
        p = para_with_style(self.doc, "1")
        add_field(p, f' TOC \\o "1-{depth}" \\h \\z \\u ')
        br = para_with_style(self.doc, "1")
        r = OxmlElement("w:r")
        r.append(el("w:br", type="page"))
        br._p.append(r)

    def title(self, text):
        para_with_style(self.doc, self.styles["doc_title"]).add_run(text)

    def heading(self, level, text):
        role = f"h{level}"
        outline = None
        if role == "h1" and self.c["rules"].get("outline_lvl_fix_h1"):
            outline = self.c["numbering"]["heading_levels"]["h1"]
        p = para_with_style(self.doc, self.styles[role], outline_lvl=outline)
        p.add_run(text)

    def body_text(self, text):
        style_id = self.styles["body"]
        facts = FACT_RE.findall(text)
        self.facts_cited.extend(facts)
        text = FACT_RE.sub("", text)
        p = para_with_style(self.doc, style_id)
        self._inline(p, text)

    def bullet(self, text):
        p = para_with_style(self.doc, self.styles["bullet"])
        numpr = OxmlElement("w:numPr")
        numpr.append(el("w:ilvl", val="0"))
        numpr.append(el("w:numId", val=self.c["numbering"]["bullet_numId"]))
        p._p.find(qn("w:pPr")).append(numpr)
        self._inline(p, text)

    def figure(self, slug, img_path):
        from docx.image.image import Image as DocxImage
        path = Path(img_path)
        if not path.is_absolute():
            # 相对路径按项目根目录解析（content.md 所在目录），仓库自带示例可直接复用
            path = self.project_dir / path
        if not path.exists():
            raise WMError(f"figure {slug} 图片不存在: {path}")
        p = para_with_style(self.doc, self.styles["figure"])
        run = p.add_run()
        try:
            with open(path, "rb") as fimg:
                img = DocxImage.from_file(fimg)
            width = Emu(914400 * img.px_width // img.horz_dpi)
            run.add_picture(str(path), width=min(width, MAX_IMG_WIDTH))
        except Exception:
            run.add_picture(str(path), width=MAX_IMG_WIDTH)

    def caption(self, kind, slug, text):
        if kind == "fig":
            style_id = self.styles["fig_caption"]
            name = f"fig_{slug}"
        else:
            style_id = self.styles["table_caption"]
            name = f"tbl_{slug}"
        p = para_with_style(self.doc, style_id)
        p.add_run(text)
        self.bm.wrap(p, name)

    def table(self, rows):
        if not rows:
            return
        cols = max(len(r) for r in rows)
        tbl = OxmlElement("w:tbl")
        tblpr = OxmlElement("w:tblPr")
        tblpr.append(el("w:tblStyle", val=self.styles["table"]))
        tblpr.append(el("w:tblW", w="5000", type="pct"))
        tbl.append(tblpr)
        grid = OxmlElement("w:tblGrid")
        for _ in range(cols):
            grid.append(el("w:gridCol", w=str(9070 // cols)))
        tbl.append(grid)
        for row in rows:
            tr = OxmlElement("w:tr")
            for ci in range(cols):
                tc = OxmlElement("w:tc")
                tcpr = OxmlElement("w:tcPr")
                tcpr.append(el("w:tcW", w=str(9070 // cols), type="dxa"))
                tc.append(tcpr)
                p = OxmlElement("w:p")
                ppr = OxmlElement("w:pPr")
                ppr.append(el("w:pStyle", val=self.styles["table_cell"]))
                p.append(ppr)
                text = row[ci] if ci < len(row) else ""
                if text:
                    r = OxmlElement("w:r")
                    t = OxmlElement("w:t")
                    t.text = text
                    r.append(t)
                    p.append(r)
                tc.append(p)
                tr.append(tc)
            tbl.append(tr)
        body = self.doc.element.body
        body.insert(list(body).index(body.find(qn("w:sectPr"))), tbl)
        # empty spacing paragraph after table (tables cannot be adjacent)
        para_with_style(self.doc, self.styles["body"])

    def _inline(self, p, text):
        """Inline renderer: **bold** *italic* {{ref:...}}."""
        pos = 0
        for m in REF_RE.finditer(text):
            if m.start() > pos:
                self._emphasis(p, text[pos:m.start()])
            kind, slug = m.group(1), m.group(2)
            name = ("fig_" if kind == "fig" else "tbl_") + slug
            add_field(p, f" REF {name} \\n \\h ")
            pos = m.end()
        if pos < len(text):
            self._emphasis(p, text[pos:])

    @staticmethod
    def _emphasis(p, text):
        # split by **bold** first, then *italic* inside plain segments
        for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
            bold = i % 2 == 1
            for j, sub in enumerate(re.split(r"\*(.+?)\*", seg)):
                if sub:
                    add_text_run(p, sub, bold=bold, italic=(j % 2 == 1) and not bold)

    def sources(self, facts):
        cited = sorted(set(self.facts_cited))
        known = {f.get("id") for f in facts.get("facts", [])}
        missing = [f for f in cited if f not in known]
        if missing:
            print(f"WARNING: content 引用了 facts.json 中不存在的事实: {', '.join(missing)}")
        if not cited:
            return
        para_with_style(self.doc, self.styles.get("sources_heading", self.styles["toc_heading"])).add_run("来  源")
        by_id = {f.get("id"): f for f in facts.get("facts", [])}
        for fid in cited:
            f = by_id.get(fid)
            if f is None:
                p = para_with_style(self.doc, self.styles["body"])
                p.add_run(f"{fid}（溯源信息缺失——需补 facts.json）")
                continue
            parts = [fid, f.get("claim", ""), f.get("url", ""),
                     f"（{f.get('date', '')}）" if f.get("date") else ""]
            p = para_with_style(self.doc, self.styles["body"])
            p.add_run(" ".join(x for x in parts if x))


def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def cmd_build(args):
    meta, ppath = load_project(resolve_project(args.project_path))
    template = meta["template"]
    contract = load_contract(template)
    pkg = TEMPLATES / template
    shell = pkg / contract.get("built_docx", "template.docx")
    if not shell.exists():
        raise WMError(f"模板壳不存在: {shell}（先完成 create-template Step 4）")
    content_md = ppath / "content.md"
    if not content_md.exists():
        raise WMError(f"缺少 {content_md}")
    facts = json.loads((ppath / "facts.json").read_text(encoding="utf-8"))

    doc = Document(str(shell))
    clear_body(doc)
    bm = Bookmarks()
    r = Renderer(doc, contract, bm, project_dir=ppath)
    r.toc_page(int(contract["rules"].get("toc_depth", 3)))

    items = parse_content(content_md.read_text(encoding="utf-8"))
    has_title = any(k == "title" for k, *_ in items)
    if not has_title:
        print("WARNING: content.md 无 @title，文档将直接以章节开始")
    for item in items:
        kind = item[0]
        if kind == "title":
            r.title(item[1])
        elif kind == "heading":
            if item[1] > int(contract["rules"].get("max_heading_level", 7)):
                raise WMError(f"标题级别超过模板上限: {'#' * item[1]}")
            r.heading(item[1], item[2])
        elif kind == "para":
            r.body_text(item[1])
        elif kind == "bullet":
            r.bullet(item[1])
        elif kind == "figure":
            r.figure(item[1], item[2])
        elif kind == "caption_fig":
            r.caption("fig", item[1], item[2])
        elif kind == "caption_tbl":
            r.caption("tbl", item[1], item[2])
        elif kind == "table":
            r.table(item[1])

    r.sources(facts)

    ts = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    out = ppath / "exports" / f"{meta['slug']}_{ts}.docx"
    out.parent.mkdir(exist_ok=True)
    doc.save(str(out))

    meta["status"] = "built"
    meta["timeline"].append({"ts": dt.datetime.now().isoformat(timespec="seconds"),
                             "event": "build", "output": out.name,
                             "facts_cited": sorted(set(r.facts_cited))})
    save_project(meta, ppath)
    print(f"[BUILD] {out.resolve()}")
    print(f"  标题/段落/图/表: 渲染完成；引用事实 {len(set(r.facts_cited))} 条")
    print(f"  下一步: doc_checker.py check {meta['slug']}")


def main():
    ap = argparse.ArgumentParser(description="word-master 构建器")
    sub = ap.add_subparsers(dest="cmd", required=True)
    sp = sub.add_parser("build")
    sp.add_argument("project_path")
    sp.set_defaults(fn=cmd_build)
    args = ap.parse_args()
    try:
        args.fn(args)
    except WMError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        sys.exit(2)


if __name__ == "__main__":
    main()
