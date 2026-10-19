# -*- coding: utf-8 -*-
"""word-master project lifecycle: init / import-sources / validate / list.

Usage:
  project_manager.py init <name> [--template qiji_heading]
  project_manager.py import-sources <project_path> <sources...>
  project_manager.py validate <project_path>
  project_manager.py list
"""
import argparse
import datetime as dt
import json
import shutil
import sys
from pathlib import Path

from wm_common import (
    CONTENT_SKELETON,
    PROJECTS,
    TEMPLATES,
    WMError,
    empty_facts,
    load_contract,
    load_project,
    resolve_project,
    save_project,
    slugify,
)

TODAY = dt.date.today().strftime("%Y%m%d")


def cmd_init(args):
    slug = f"{slugify(args.name)}_{TODAY}"
    p = PROJECTS / slug
    if p.exists():
        raise WMError(f"项目已存在: {p}")
    template = args.template
    load_contract(template)  # fail fast if template not installed
    (p / "sources" / "converted").mkdir(parents=True)
    (p / "exports").mkdir()
    meta = {
        "$schema": "word-master.project.v1",
        "name": args.name,
        "slug": slug,
        "template": template,
        "created": dt.datetime.now().isoformat(timespec="seconds"),
        "status": "initialized",
        "sources": [],
        "timeline": [],
    }
    save_project(meta, p)
    (p / "facts.json").write_text(json.dumps(empty_facts(), ensure_ascii=False, indent=2),
                                  encoding="utf-8")
    (p / "content.md").write_text(CONTENT_SKELETON.format(template=template), encoding="utf-8")
    print(f"项目已创建: {p.resolve()}")
    print(f"  模板: {template}  状态: initialized")
    print("  下一步: import-sources → 确认大纲 → 编写 content.md")


def cmd_import(args):
    meta, p = load_project(resolve_project(args.project_path))
    conv = p / "sources" / "converted"
    for s in args.sources:
        src = Path(s)
        if not src.exists():
            raise WMError(f"来源不存在: {src}")
        dst = p / "sources" / src.name
        shutil.copy2(src, dst)
        entry = {"file": src.name, "copied_from": str(src.resolve()),
                 "size": src.stat().st_size}
        if src.suffix.lower() == ".docx":
            txt = convert_docx(src, conv / f"{src.stem}.txt")
            entry["converted_text"] = str(txt.relative_to(p))
        elif src.suffix.lower() == ".pdf":
            entry["converted_text"] = None
            entry["note"] = "PDF 转文本需要 PyMuPDF（MVP 仅复制原文）"
        meta["sources"].append(entry)
        print(f"已导入: {src.name}" + (f" → {entry['converted_text']}" if entry.get("converted_text") else ""))
    meta["status"] = "sources_imported" if meta["status"] == "initialized" else meta["status"]
    meta["timeline"].append({"ts": dt.datetime.now().isoformat(timespec="seconds"),
                             "event": "import-sources", "count": len(args.sources)})
    save_project(meta, p)


def convert_docx(src, out_txt):
    from docx import Document
    doc = Document(src)
    lines = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            style = para.style.name if para.style is not None else ""
            lines.append(f"[{style}] {t}" if "heading" in style.lower() else t)
    for i, tbl in enumerate(doc.tables, 1):
        lines.append(f"--- 表格 {i} ---")
        for row in tbl.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    out_txt.parent.mkdir(parents=True, exist_ok=True)
    out_txt.write_text("\n".join(lines), encoding="utf-8")
    return out_txt


def cmd_validate(args):
    meta, p = load_project(resolve_project(args.project_path))
    problems = []
    if meta.get("template") and not (TEMPLATES / meta["template"] / "contract.json").exists():
        problems.append(f"模板未入库: {meta['template']}")
    for f in ("facts.json", "content.md"):
        fp = p / f
        if not fp.exists():
            problems.append(f"缺少 {f}")
        elif f == "facts.json":
            try:
                json.loads(fp.read_text(encoding="utf-8"))
            except json.JSONDecodeError as e:
                problems.append(f"facts.json 不是合法 JSON: {e}")
    if problems:
        print("INVALID")
        for x in problems:
            print(f"  - {x}")
        sys.exit(1)
    print(f"VALID  {p}")
    print(f"  模板={meta['template']} 状态={meta['status']} 来源={len(meta.get('sources', []))}条")


def cmd_list(_args):
    if not PROJECTS.exists() or not any(PROJECTS.iterdir()):
        print("（无项目）")
        return
    for d in sorted(PROJECTS.iterdir()):
        pj = d / "project.json"
        if pj.exists():
            m = json.loads(pj.read_text(encoding="utf-8"))
            print(f"{d.name}  模板={m.get('template')}  状态={m.get('status')}")


def main():
    ap = argparse.ArgumentParser(description="word-master 项目管理")
    sub = ap.add_subparsers(dest="cmd", required=True)

    sp = sub.add_parser("init")
    sp.add_argument("name")
    sp.add_argument("--template", default="qiji_heading")
    sp.set_defaults(fn=cmd_init)

    sp = sub.add_parser("import-sources")
    sp.add_argument("project_path")
    sp.add_argument("sources", nargs="+")
    sp.set_defaults(fn=cmd_import)

    sp = sub.add_parser("validate")
    sp.add_argument("project_path")
    sp.set_defaults(fn=cmd_validate)

    sp = sub.add_parser("list")
    sp.set_defaults(fn=cmd_list)

    args = ap.parse_args()
    try:
        args.fn(args)
    except WMError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        sys.exit(2)


if __name__ == "__main__":
    main()
